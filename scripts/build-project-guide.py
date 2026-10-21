"""Build the Chinese project handbook with the Codex bundled Python runtime.

Usage: python scripts/build-project-guide.py --ui-verification "Verified UI facts"
Render the DOCX with the documents skill's render_docx.py before delivery.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BLACK = "000000"
FONT = "Microsoft YaHei"
UI_VERIFICATION = "浏览器已跑通模拟语音字幕、任务确认、云端执行、结果编辑、追问和人工入库，重启后数据保留。本地新任务也完成选句、显示原文及确认入库，公式和不等号原样保留。项目保留标明示例的试运行记录，用来展示流程，不代表真实实验。桌面 1366×900 与 360 像素窄布局已检查，模式切换和连接检查通过。"


def font_for(style, size: float, bold: bool = False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor.from_string(BLACK)
    props = style.element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.append(fonts)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{name}"), FONT)
    for node in list(props):
        if node.tag == qn("w:color"):
            node.attrib.pop(qn("w:themeColor"), None)
    paragraph_props = style.element.find(qn("w:pPr"))
    if paragraph_props is not None:
        for border in list(paragraph_props.findall(qn("w:pBdr"))):
            paragraph_props.remove(border)


def paragraph(doc, text: str, bold_lead: str | None = None, after: int = 6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def heading(doc, text: str, level: int = 2):
    return doc.add_heading(text, level=level)


def step(doc, number: int, text: str):
    return paragraph(doc, f"{number}  {text}", after=7)


def code(doc, value: str):
    p = doc.add_paragraph(style="Code")
    p.add_run(value)
    return p


def table(doc, headers: list[str], rows: list[tuple[str, ...]], widths: list[float]):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for index, width in enumerate(widths):
        tbl.columns[index].width = Inches(width)
    props = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", "95"), ("bottom", "95"), ("left", "105"), ("right", "105")):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    props.append(margins)
    for index, label in enumerate(headers):
        tbl.rows[0].cells[index].text = label
    for values in rows:
        cells = tbl.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    for row_index, row in enumerate(tbl.rows):
        tr_props = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        tr_props.append(no_split)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            tr_props.append(repeat)
        for col_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E8EDF1" if row_index == 0 else ("F7F8FA" if row_index % 2 == 0 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                if col_index == 0 and len(headers) > 2:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(10.5)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor.from_string(BLACK)
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return tbl


def hyperlink(doc, label: str, url: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.add_run(f"{label}  ")
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLACK)
    props.append(color)
    run.append(props)
    value = OxmlElement("w:t")
    value.text = url
    run.append(value)
    element.append(run)
    p._p.append(element)
    return p


def new_page(doc, title: str):
    p = heading(doc, title, 1)
    p.paragraph_format.page_break_before = True


def build(ui_verification: str, destination: Path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.footer_distance = Inches(0.28)

    for name, size, bold in (("Normal", 11, False), ("Title", 23, True), ("Subtitle", 11, False), ("Heading 1", 16, True), ("Heading 2", 12.5, True), ("Heading 3", 11, True)):
        font_for(doc.styles[name], size, bold)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.20
    normal.space_after = Pt(6)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        fmt = doc.styles[name].paragraph_format
        fmt.space_before = Pt(10 if name != "Heading 1" else 0)
        fmt.space_after = Pt(7 if name != "Heading 1" else 12)
        fmt.keep_with_next = True
    doc.styles["Title"].paragraph_format.space_after = Pt(7)
    doc.styles["Title"].paragraph_format.keep_with_next = True
    code_style = doc.styles.add_style("Code", 1)
    font_for(code_style, 10)
    code_style.paragraph_format.space_after = Pt(5)
    code_style.paragraph_format.line_spacing = 1.1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    for run in footer.runs:
        run.font.name = FONT
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(BLACK)

    doc.core_properties.title = "Personal Agent 项目说明"
    doc.core_properties.subject = "电脑端科研学习助手的使用与实现"
    doc.core_properties.author = "Personal Agent"
    doc.core_properties.keywords = "科研学习 任务队列 知识库 Windows Ollama"

    doc.add_paragraph("Personal Agent 项目说明", style="Title")
    doc.add_paragraph("电脑端科研学习助手使用与实现手册", style="Subtitle")
    paragraph(doc, "2026 年 10 月 20 日", after=14)
    heading(doc, "1 项目目的与当前状态", 1)
    paragraph(doc, "这个项目把课堂笔记、论文摘录和讨论记录变成可以跟进的学习或研究任务。你可从 OpenDCAI/ThinkFlow 读取资料，确认目标后交给后台整理，再审阅、追问，最后手动确认写回原笔记本。也可粘贴材料并存入本地库。现在已在 Windows 电脑跑通，整个流程不依赖树莓派。")
    paragraph(doc, "树莓派在原方案中承担小型设备入口的作用。先在电脑上开发，可以验证任务流程和软件结构；以后再补麦克风、音箱、实体按键及设备适配。目前的运行与速度结果都来自 Windows 电脑，不能代表树莓派性能。")
    heading(doc, "从一段笔记到一条知识")
    paragraph(doc, "例如，从 ThinkFlow 选取两份组会资料，填写“整理待验证问题和下一步实验”，预览并确认后开始处理。生成后，你可以追问“材料里还有哪些未确定的条件”，对照原文修正，再确认写回。资料和结果可以在 ThinkFlow 原生界面中查看。")
    paragraph(doc, "实际流程为：输入目标与材料 → 预览并确认 → 排队处理 → 查看结果并追问 → 人工确认入库。保留草稿不会触发处理；没有正文时，程序只列出需要补充的信息。")
    heading(doc, "当前电脑的使用方式")
    paragraph(doc, "这台电脑没有麦克风，当前界面使用模拟语音、字幕和文字录入演示入口。模拟语音把示例文本带入任务表单，不录音；当前产品没有真实语音识别 ASR 链路。浏览器朗读属于语音输出，与识别不是同一功能。")
    paragraph(doc, "程序支持三种手动选择的处理方式：本地规则提取、本机 Ollama 模型，以及通过本机代理调用云端模型。检测到已有代理配置时，首次默认采用 gpt-6-luna；你保存过的选择会保留。模式之间没有自动切换或自动云端兜底。")
    paragraph(doc, "可先按第 2 节跑通任务，再用第 3 节选择处理方式。第 4 至第 6 节解释数据、验证及设备迁移；第 7 节说明 ThinkFlow 的使用、能力与故障恢复。")

    new_page(doc, "2 启动与一次完整使用")
    heading(doc, "启动和停止")
    paragraph(doc, "项目目录位于：")
    code(doc, r"C:\Users\19484\Documents\ChatGPT\personal agent")
    paragraph(doc, "双击 start.cmd。启动器会复用已有服务，或隐藏启动后端、已安装的 Ollama 和 ThinkFlow，再打开 http://127.0.0.1:4317。ThinkFlow 原生界面在 8437 端口。需要 Node.js 22 或更新版本，本项目 Node 后端无需 npm 依赖。")
    code(doc, "powershell.exe -NoProfile -File .\\scripts\\start.ps1")
    paragraph(doc, "关闭网页不会停止后台服务。结束使用时，执行下列命令停止本项目登记并核对过身份的进程；其他已有服务不会被按端口直接关闭。")
    code(doc, "powershell.exe -NoProfile -File .\\scripts\\stop.ps1")
    paragraph(doc, "启动器不会更改 Windows 脚本安全策略。若执行策略阻止脚本，可在项目目录运行 node server.mjs 启动网页后端，再使用已有模型服务，或在界面选择规则提取。")
    heading(doc, "完成第一项任务")
    step(doc, 1, "进入工作台，选择学习整理、研究问题整理或资料整理。可使用示例或模拟语音，也可以直接输入。")
    step(doc, 2, "填写目标，点击“选取资料”从 ThinkFlow 读取预览并带入正文，或直接粘贴材料。手工登记的普通 URL 只保留出处，不会自动打开网页。")
    step(doc, 3, "预览任务，核对目标与材料，再点击“开始处理”。“保留草稿”只保存，不调用模型。")
    step(doc, 4, "在“我的任务”查看排队、处理中或完成状态。后端按顺序处理任务，页面关闭后仍可回来查看。")
    step(doc, 5, "对照原文检查结果，必要时编辑或追问。朗读采用浏览器语音合成，可点击“停止”中止播报。")
    step(doc, 6, "核对后，ThinkFlow 任务点击“确认并写回 ThinkFlow”，普通任务点击“确认存入知识库”。知识库支持搜索、Markdown 下载和本地业务数据 JSON 导出。")
    heading(doc, "建议用来试跑的材料")
    paragraph(doc, "目标：整理识别实验的已知事实与下一步。正文：安静环境中 10 次录音有 9 次识别正确，噪声环境中 10 次有 6 次正确。下一步固定麦克风距离后重复测试。目前样本较少，不能推断普遍表现。")
    paragraph(doc, "这段材料可检查程序能否保留样本数量和条件。生成内容如果省略限制条件，先修正再入库。")

    new_page(doc, "3 三种处理方式与数据去向")
    paragraph(doc, "在“设置”选择处理方式，点击“保存设置”，再点击“测试已保存的配置”。填入预设后仍需要保存。连接测试只读取模型列表并核对所选模型，真实生成能力要通过提交任务验证。")
    table(doc, ["方式", "地址和模型", "材料去向与能力"], [
        ("本地规则提取", "无需模型服务", "由程序在本机抽取句子，标记问题和行动候选；追问是关键词检索，不是 AI 总结或推理。"),
        ("Ollama 本机模型", "127.0.0.1:11434\nqwen2.5:1.5b", "模型辅助选择相关原句，程序核对编号后呈现原文。追问也检索原句，不用于科学推理。"),
        ("本机代理调用云端", "127.0.0.1:8317/v1\ngpt-6-luna", "请求通过已有 OpenAI 兼容代理发往配置的模型服务。入口地址在本机，推理仍发生在云端。"),
    ], [1.15, 2.0, 3.85])
    heading(doc, "本地模型的约束")
    paragraph(doc, "项目内已部署 Ollama 0.40.6 和 qwen2.5:1.5b。模型主要文件约 986 MB，采用 Q4_K_M 量化。程序为 Ollama 设置 8,192 token 上下文，最多生成 400 token 的选句结果；程序只采用经校验的原文编号，不采用小模型改写的正文。token 是模型内部的文本单位，与字符数不同。")
    paragraph(doc, "本地正文最多 5,000 字符，系统要求与完整任务数据合计最多 6,500 字符。追问同时包含已有结果和最近 8 条历史消息，每条完整保留；超限需拆分。模型只返回原句编号，结构错误、没有有效编号或达到输出上限会明确报错，不把自由生成的文字当正文。问题和行动候选仍由规则分类。")
    paragraph(doc, "需要重新安装、续传模型或验证本地生成时，执行安装脚本。启动器不会自动下载大文件。")
    code(doc, "powershell.exe -NoProfile -File .\\scripts\\setup-local-model.ps1")
    heading(doc, "材料与密钥如何处理")
    paragraph(doc, "规则模式不调用模型接口。Ollama 模式把正文交给所选本机服务；云端模式则可能上传正文、任务要求和追问上下文。请按材料本身的使用要求选择模式。")
    paragraph(doc, "手动填写的 API Key 只留在当前后端进程内存，不写入业务数据或导出；重启后需重新提供。已有本机代理的客户端凭证按需读取到内存，项目不复制 OAuth 文件。")

    new_page(doc, "4 数据流与后端架构")
    paragraph(doc, "浏览器负责表单、任务列表、结果显示和播报。Node.js 后端只监听本机地址 127.0.0.1:4317，负责校验请求、保存数据和串行执行队列。处理模块依据已保存的模式调用规则、Ollama 或 OpenAI 兼容接口，再把结果交回后端。")
    paragraph(doc, "数据流为：读取资料或粘贴正文 → 保存草稿 → 确认目标 → 队列处理 → 审阅结果 → 确认保存。ThinkFlow 任务会创建原生文档，读回核对后再记本地确认条目。模型调用失败会报错，不自动换成其他模式。")
    heading(doc, "任务状态和异步的含义")
    paragraph(doc, "任务从草稿进入队列，再变为处理中，随后完成、失败或取消。网页只是操作入口；后端运行期间，关闭页面不会中断处理。这里的异步是本机后台队列，并没有额外部署一直在线的云端任务服务。")
    paragraph(doc, "电脑睡眠、关机或后端退出后，任务不能继续运行。服务重启会把中断的任务标为失败，保留信息供你检查，并允许重新确认执行。这避免把未完成工作显示为已完成。")
    heading(doc, "需要知道的文件")
    table(doc, ["文件或目录", "作用"], [
        ("server.mjs", "本机 HTTP 服务、任务状态机、串行队列及数据持久化。"),
        ("lib/processor.mjs", "规则提取、模型请求、追问、连接检查和长度限制。"),
        ("lib/thinkflow.mjs", "ThinkFlow 资料读取、来源快照、原生文档回写与读回校验。"),
        ("public/", "中文网页界面、模拟输入、任务管理、知识库和浏览器播报。"),
        ("scripts/", "启动、停止、本地模型部署，以及本说明的可维护生成脚本。"),
        ("data/workspace.json", "任务、结果、对话、来源、设置及确认后的知识条目。"),
        ("data/thinkflow/outputs", "ThinkFlow 原生资料、成果文档及其版本的实际存储目录。"),
    ], [2.05, 4.95])
    heading(doc, "保存和备份")
    paragraph(doc, "备份前先停止服务，再复制整个 data 目录，同时保留应用记录和 ThinkFlow 原生文件。网页 JSON 导出只包含本地业务记录，不包括 ThinkFlow 原始资料和版本目录；当前没有图形化导入恢复功能。.runtime 中的程序、模型可重新安装。")

    new_page(doc, "5 原方案对应与验证证据")
    paragraph(doc, "原图描述的是包含设备入口、任务处理和知识沉淀的完整方案。本次完成电脑端原型，软件流程与硬件验证需要分别判断。")
    table(doc, ["原方案内容", "当前实现与边界"], [
        ("轻量任务入口", "文字、示例和模拟语音可建任务；无唤醒词或全局热键。"),
        ("语音与设备原型", "仅模拟输入与浏览器播报；无真实 ASR 链路。树莓派未部署。"),
        ("本地优先与云端补充", "三种方式手动切换；没有自动路由或云端兜底。"),
        ("异步总结与追问", "三类任务可后台排队、整理及追问；电脑需保持运行。"),
        ("知识沉淀", "已接通 ThinkFlow 原始资料读取与确认后原生文档回写；无研究关系图。"),
        ("六类场景验证", "硬件六类场景验证尚未完成，软件测试不能替代。"),
    ], [1.65, 5.35])
    heading(doc, "已经执行的检查")
    paragraph(doc, "自动测试共 30 项，均通过：处理模块 17 项、后端 7 项、ThinkFlow 接入 6 项。新增覆盖资料归属、预览变化、人工确认门槛、并发确认、读回不一致、响应丢失及重启恢复。模拟服务用于验证契约，不作为真实模型能力的证据。")
    paragraph(doc, "已实际跑通复用服务、新建后台服务、停止登记服务及再次启动。停止操作保留了原有代理；进程身份保护的异常分支仅做静态审查。")
    if ui_verification:
        paragraph(doc, ui_verification)
    paragraph(doc, "云端 gpt-6-luna 已完成真实示例生成，单次约 5.93 秒。实际测试发现 1.5B 模型会错改公式，因此本地仅辅助选句和资料线索检索，复杂归纳与推理交给云端。原句可能遗漏上下文，云端也可能推断错误；入库前核对数字、条件与引用。耗时只对应当前 Windows 电脑的一次观察，不是树莓派基准测试。")

    new_page(doc, "6 树莓派迁移与日常排错")
    heading(doc, "以后如何迁移到树莓派")
    paragraph(doc, "核心任务流程使用 Node.js、HTTP 接口和本地 JSON，迁移时可以保留这些思路。先在树莓派安装合适的 Linux 与 Node.js，复制程序和数据，修改 Windows 专用的启动及路径配置，再验证网页与队列。现有 Windows 便携 Ollama 可执行文件不能直接复制到树莓派运行。")
    paragraph(doc, "再接入麦克风、音箱和按键，逐项测试音频设备、权限、识别延迟、回声及播报打断。可让树莓派只负责采集和反馈，把较重推理留在电脑或云端；也可评估合适的 ARM 本地模型。是否能在设备上达到可用速度，必须按实际内存、散热与模型重新测量。")
    heading(doc, "常见问题的检查顺序")
    paragraph(doc, "网页打不开：先运行 start.cmd，再检查 4317 端口是否被其他服务占用。查看 .runtime/server.stderr.log。不要仅凭端口号结束不认识的进程。")
    paragraph(doc, "模型连接失败：核对保存的模式、地址和模型名。本地模型查看 .runtime/ollama.stderr.log；代理查看 .runtime/proxy.stderr.log。连接测试未找到模型时，先确认安装或账户权限。")
    paragraph(doc, "任务失败或内容不完整：检查错误信息；本地模式超长、选句无效或达到输出上限时，缩小材料与追问范围后重试。后端重启后的中断任务需要重新确认执行。云端结果也需核对原文。")
    paragraph(doc, "朗读没有声音：检查浏览器支持、音量和可用中文声音。模拟输入本身不会录音；当前没有真实 ASR 链路，不能用它测试识别效果。")
    paragraph(doc, "需要重新验证代码：在项目目录运行 npm test。需要备份成果：在知识库导出 JSON；需要复用说明：修改 scripts/build-project-guide.py 后重新生成并检查排版。")
    heading(doc, "继续阅读的官方资料")
    hyperlink(doc, "Ollama Chat API", "https://docs.ollama.com/api/chat")
    hyperlink(doc, "OpenAI Chat Completions", "https://platform.openai.com/docs/api-reference/chat/create")
    hyperlink(doc, "树莓派入门文档", "https://www.raspberrypi.com/documentation/computers/getting-started.html")

    new_page(doc, "7 ThinkFlow 资料与结果管理")
    paragraph(doc, "ThinkFlow 现在承载原始资料与经确认的成果；Personal Agent 负责目标确认、任务队列和模型处理。部署使用 OpenDCAI/ThinkFlow 固定提交 1eecae5，原版后端和前端均已在本机运行。")
    heading(doc, "怎样从资料得到正式成果")
    step(doc, 1, "在 4317 工作台点击“选取资料”。选择已有笔记本；也可展开弹窗下方的选项，新建笔记本并保存文字资料。保存原始资料不会触发模型。")
    step(doc, 2, "勾选材料后“读取预览”，核对正文，再点击“使用已读取资料”。当前一次最多 20 份，正文合计含标记最多 100,000 字符。Ollama 仍受更小预算限制。")
    step(doc, 3, "确认目标后运行。结果完成时先审阅，可编辑并保存。资料读取、生成和追问都不会自动创建 ThinkFlow 成果文档。")
    step(doc, 4, "点击“确认并写回 ThinkFlow”。写入并读取核对成功后，可点“读取 ThinkFlow 文档”查看实时正文，或在原生工作台打开笔记本并点击右侧文档标题。")
    heading(doc, "来源记录和失败恢复")
    paragraph(doc, "任务保留正文快照、资料 ID、SHA-256 和读取时间；新任务还记录各来源句段范围，新回写文档保存原文快照。SHA-256 用于比较内容是否相同，不证明模型结论正确。普通 URL 仍然只作登记。")
    paragraph(doc, "第一次确认后锁定该次内容。若响应丢失或服务中断，恢复后点击“重试核对并写回”，程序先查任务对应文档，核对一致后恢复，不重复创建。若原文档已被外部修改，会报冲突并保留，不自动覆盖。")
    heading(doc, "已经接通的范围")
    paragraph(doc, "真实联调已完成两份 Markdown 资料读取、gpt-6-luna 整理、编辑、确认回写与原生 UI 查看。确认前结果文档为 0，确认及重复确认后仍为 1；读回正文与审阅稿一致。资料中的实验计划是示例，不代表已完成实验。")
    paragraph(doc, "原生素材的“待入库”指未建立向量索引，正文已保存且可以直接读取。当前没有启用 embedding 和文档 RAG，也未接通 BM25 融合重排、Skills 或 Plan-and-Execute。PDF 读取接口允许文本层 PDF，尚未单独实测。")
    paragraph(doc, "新电脑先运行 scripts/setup-thinkflow.ps1，需 Git、uv、Node.js/npm 和网络。常规启动用 start.cmd；排错看 .runtime/thinkflow.stderr.log。详情见 docs/ThinkFlow接入说明.md。")

    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    print(destination)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build the Chinese Personal Agent handbook")
    parser.add_argument("--output", type=Path, default=ROOT / "docs" / "Personal Agent 项目说明.docx")
    parser.add_argument("--ui-verification", default=UI_VERIFICATION, help="Confirmed browser UI validation facts to include")
    args = parser.parse_args()
    build(args.ui_verification, args.output.resolve())
